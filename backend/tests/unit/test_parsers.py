from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.ai.parsers import (
    DocxParser,
    ParseError,
    PdfParser,
    TxtParser,
)


def make_pdf_with_text(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


def make_blank_pdf(encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if encrypted:
        writer.encrypt("secret-password")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(
    paragraphs: tuple[str, ...] = (), table_rows: tuple[tuple[str, ...], ...] = ()
) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, cell_text in enumerate(row):
                table.cell(row_index, col_index).text = cell_text
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_extracts_text() -> None:
    result = PdfParser().parse(make_pdf_with_text("Quarterly planning sync 2026"), "plan.pdf")
    assert result.parser == "pdf"
    assert result.file_name == "plan.pdf"
    assert "Quarterly planning sync 2026" in result.content
    assert result.metadata["page_count"] == 1
    assert result.metadata["char_count"] == len(result.content)


def test_pdf_without_text_layer_reports_reason() -> None:
    with pytest.raises(ParseError) as excinfo:
        PdfParser().parse(make_blank_pdf(), "scan.pdf")
    assert excinfo.value.reason == "no_text_layer"
    assert excinfo.value.error_code == "FILE_PARSE_FAILED"
    assert "OCR" in excinfo.value.message


def test_pdf_encrypted_reports_reason() -> None:
    with pytest.raises(ParseError) as excinfo:
        PdfParser().parse(make_blank_pdf(encrypted=True), "locked.pdf")
    assert excinfo.value.reason == "encrypted_pdf"


def test_pdf_corrupt_reports_invalid_file() -> None:
    with pytest.raises(ParseError) as excinfo:
        PdfParser().parse(b"this is not a pdf at all", "broken.pdf")
    assert excinfo.value.reason == "invalid_file"


def test_docx_extracts_paragraphs_and_tables() -> None:
    result = DocxParser().parse(
        make_docx(
            paragraphs=("Agenda item one", "Decision: ship on Friday"),
            table_rows=(("Owner", "Due"), ("Alice", "2026-10-10")),
        ),
        "minutes.docx",
    )
    assert result.parser == "docx"
    assert "Agenda item one" in result.content
    assert "Decision: ship on Friday" in result.content
    assert "Owner\tDue" in result.content
    assert "Alice\t2026-10-10" in result.content
    assert result.metadata["paragraph_count"] == 2
    assert result.metadata["table_count"] == 1
    assert result.metadata["char_count"] == len(result.content)


def test_docx_without_content_reports_empty() -> None:
    with pytest.raises(ParseError) as excinfo:
        DocxParser().parse(make_docx(), "empty.docx")
    assert excinfo.value.reason == "empty_content"


def test_docx_corrupt_reports_invalid_file() -> None:
    with pytest.raises(ParseError) as excinfo:
        DocxParser().parse(b"garbage-bytes", "broken.docx")
    assert excinfo.value.reason == "invalid_file"


def test_txt_utf8() -> None:
    result = TxtParser().parse("会议纪要\n第二行".encode(), "notes.txt")
    assert result.content == "会议纪要\n第二行"
    assert result.metadata["encoding"] == "utf-8"
    assert result.metadata["char_count"] == len(result.content)


def test_txt_utf8_bom() -> None:
    result = TxtParser().parse("BOM header".encode("utf-8-sig"), "bom.txt")
    assert result.content == "BOM header"
    assert result.metadata["encoding"] == "utf-8-sig"


def test_txt_gb18030_fallback() -> None:
    original = "季度规划会议"
    result = TxtParser().parse(original.encode("gb18030"), "gbk.txt")
    assert result.content == original
    assert result.metadata["encoding"] == "gb18030"


def test_txt_normalizes_line_endings() -> None:
    result = TxtParser().parse(b"line1\r\nline2\rline3", "win.txt")
    assert result.content == "line1\nline2\nline3"


def test_txt_empty_reports_reason() -> None:
    with pytest.raises(ParseError) as excinfo:
        TxtParser().parse(b"   \n  ", "blank.txt")
    assert excinfo.value.reason == "empty_content"


def test_txt_binary_reports_invalid_file() -> None:
    with pytest.raises(ParseError) as excinfo:
        TxtParser().parse(b"\x89PNG\r\n\x1a\n\x00\x00", "fake.txt")
    assert excinfo.value.reason == "invalid_file"


def test_parsed_document_as_result_matches_contract() -> None:
    result = TxtParser().parse(b"hello", "note.txt").as_result("material_001")
    assert result == {
        "material_id": "material_001",
        "status": "processed",
        "file_name": "note.txt",
        "content": "hello",
        "metadata": {"parser": "txt", "encoding": "utf-8", "char_count": 5},
    }
