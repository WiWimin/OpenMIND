from __future__ import annotations

from io import BytesIO

from docx import Document

from app.ai.parsers.base import BaseParser, ParsedDocument, ParseError


class DocxParser(BaseParser):
    parser_name = "docx"

    def parse(self, data: bytes, file_name: str) -> ParsedDocument:
        try:
            document = Document(BytesIO(data))
        except Exception as exc:
            raise ParseError("invalid_file", f"DOCX 文件无法读取或已损坏：{file_name}") from exc

        parts: list[str] = [paragraph.text.strip() for paragraph in document.paragraphs]
        table_count = 0
        for table in document.tables:
            table_count += 1
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))

        content = "\n".join(part for part in parts if part)
        if not content:
            raise ParseError("empty_content", f"DOCX 文件未包含任何文本：{file_name}")

        return ParsedDocument(
            file_name=file_name,
            content=content,
            parser=self.parser_name,
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": table_count,
                "char_count": len(content),
            },
        )
